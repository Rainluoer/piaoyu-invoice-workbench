"""统计表固定分类、留空规则、同步生成、编辑持久化及配套留存。"""
import copy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import threading
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request

from docx import Document
from docx.oxml.ns import qn
import test_receipt_jobs as fixtures
from server import Handler, LocalServer, Store
from statistics_table import rows_from_invoices, validate_rows, write_statistics, HEADERS


class StatisticsTableTests(unittest.TestCase):
    def test_fixed_categories_price_and_unknown_fields(self):
        invoices = [dict(total='12.34', items=[dict(category='化学试剂', product='缓冲液', spec='100ml', amount=10, tax=2.34)]),
                    dict(total=0, items=[dict(category='设备', product='电子天平', amount=100)]),
                    dict(total=20, items=[dict(product='螺丝', amount='10.10', tax='1.31'), dict(product='PLA', amount=None)])]
        rows = rows_from_invoices(invoices)
        self.assertEqual([r['category'] for r in rows], ['试剂', '仪器', '耗材', '耗材'])
        self.assertEqual([r['price'] for r in rows], ['12.34', '0.00', '11.41', ''])
        self.assertEqual(rows[0]['name'], '缓冲液 100ml')
        self.assertTrue(all(r['url'] == r['safety'] == r['purpose'] == '' for r in rows))

    def test_only_explicit_urls_and_reagent_safety_and_no_invented_categories(self):
        rows = rows_from_invoices([dict(total=None, items=[dict(product='试剂盒'), dict(product='电线')])])
        submitted = copy.deepcopy(rows)
        submitted[0].update(safety='用户填写的等级', url='https://example.com/safety')
        submitted[1].update(safety='不应输出', url='https://example.com/product', price='999.99', name='篡改名称')
        result = validate_rows(submitted, rows)
        self.assertEqual(result[0]['safety'], '用户填写的等级')
        self.assertEqual(result[1]['safety'], '')
        self.assertEqual(result[1]['url'], 'https://example.com/product')
        self.assertEqual((result[1]['name'], result[1]['price']), ('电线', ''))
        for category in ('其他', '科研材料', '', '五金'):
            invalid = copy.deepcopy(submitted)
            invalid[0]['category'] = category
            with self.assertRaisesRegex(ValueError, '只能选择'):
                validate_rows(invalid, rows)
        for url in ('javascript:alert(1)', 'file:///C:/test', 'https://user:password@example.com', 'https://example.com/a b'):
            invalid = copy.deepcopy(submitted)
            invalid[0]['url'] = url
            with self.assertRaisesRegex(ValueError, '网址'):
                validate_rows(invalid, rows)
        with self.assertRaises(ValueError):
            validate_rows([submitted[0], submitted[0]], rows)

    def test_document_preserves_schema_without_sample_rows_or_bottom_notes(self):
        rows = rows_from_invoices([dict(total=None, items=[dict(product='试剂盒'), dict(product='显微镜'), dict(product='连接线')])])
        rows[0]['safety'] = '按用户提供填写'
        rows[1]['safety'] = '不得输出'
        with TemporaryDirectory() as directory:
            path = Path(directory) / '统计表格.docx'
            write_statistics(path, rows)
            document = Document(path)
        self.assertEqual(tuple(c.text for c in document.tables[0].rows[0].cells), HEADERS)
        self.assertEqual(len(document.tables[0].rows), 4)
        self.assertEqual([r.cells[6].text for r in document.tables[0].rows[1:]], ['按用户提供填写', '', ''])
        self.assertTrue(all(r.cells[4].text == '' for r in document.tables[0].rows[1:]))
        body = list(document.element.body)
        self.assertEqual(body[-2].tag, qn('w:tbl'))
        self.assertTrue(all(not p.text for section in document.sections for p in section.footer.paragraphs))


class StatisticsJobTests(unittest.TestCase):
    def setUp(self):
        self.fixture = fixtures.ReceiptJobTests()
        self.fixture.setUp()
        self.addCleanup(self.fixture.tearDown)
        self.store = self.fixture.store

    def generate(self):
        with patch('receipt_jobs.subprocess.run', side_effect=self.fixture.fake_script):
            return self.store.receipts.create([self.fixture.row['id']], background=False)

    def test_generated_pair_download_metadata_and_five_group_retention(self):
        jobs = [self.generate() for _ in range(6)]
        for job in jobs:
            self.assertEqual(job['status'], 'done')
            self.assertEqual(job['statistics_error'], '')
        self.assertFalse(Path(jobs[0]['output']).parent.exists())
        documents = self.store.documents()
        self.assertEqual(len(documents), 5)
        self.assertTrue(all(doc['statistics'] for doc in documents))
        for doc in documents:
            table = Document(doc['statistics']['path']).tables[0]
            self.assertEqual(table.rows[1].cells[5].text, '12.34')
            self.assertEqual(table.rows[1].cells[0].text, '耗材')
        self.assertEqual(self.fixture.original.read_bytes(), b'%PDF fixture for scheduler')

    def test_saved_selection_survives_restart_and_does_not_change_receipt(self):
        job = self.generate()
        receipt_before = Path(job['output']).read_bytes()
        rows = self.store.receipts.statistics_data(job['id'])['rows']
        rows[0].update(category='试剂', purpose='用户指定用途', url='https://example.com/safety', safety='用户指定')
        self.store.receipts.save_statistics(job['id'], rows)
        self.assertEqual(Path(job['output']).read_bytes(), receipt_before)
        self.store.db.close()
        self.fixture.store = self.store = Store(self.fixture.root, self.fixture.root / 'app/data')
        restored = self.store.receipts.statistics_data(job['id'])['rows']
        self.assertEqual(restored[0]['category'], '试剂')
        self.assertEqual(restored[0]['safety'], '用户指定')
        self.assertEqual(self.store.db.execute("SELECT COUNT(*) FROM events WHERE action='修改统计表格'").fetchone()[0], 1)

    def test_generation_error_preserves_receipt_and_editor_can_retry(self):
        with patch('receipt_jobs.write_statistics', side_effect=PermissionError('file busy')):
            job = self.generate()
        self.assertEqual(job['status'], 'done')
        self.assertIn('统计表格未能生成', job['statistics_error'])
        self.assertTrue(Path(job['output']).is_file())
        self.assertIsNone(self.store.documents()[0]['statistics'])
        rows = self.store.receipts.statistics_data(job['id'])['rows']
        done = self.store.receipts.save_statistics(job['id'], rows)
        self.assertEqual(done['statistics_error'], '')
        self.assertIsNotNone(self.store.documents()[0]['statistics'])

    def test_failed_audit_rolls_back_statistics_file_and_saved_rows(self):
        job = self.generate()
        path = Path(self.store.documents()[0]['statistics']['path'])
        before = path.read_bytes()
        rows = self.store.receipts.statistics_data(job['id'])['rows']
        changed = copy.deepcopy(rows)
        changed[0]['category'] = '仪器'
        with patch.object(self.store, 'event', side_effect=RuntimeError('audit failed')):
            with self.assertRaisesRegex(RuntimeError, 'audit failed'):
                self.store.receipts.save_statistics(job['id'], changed)
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(self.store.receipts.statistics_data(job['id'])['rows'], rows)
        self.assertFalse(self.store.receipts.active.locked())

    def test_http_edit_requires_key_and_generated_statistics_can_download(self):
        job = self.generate()
        http = LocalServer(('127.0.0.1', 0), Handler)
        http.store = self.store
        thread = threading.Thread(target=http.serve_forever, daemon=True)
        thread.start()
        base = f'http://127.0.0.1:{http.server_port}'
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        route = base + '/api/receipt-jobs/' + job['id'] + '/statistics'
        try:
            with opener.open(route) as response:
                rows = json.load(response)['rows']
            body = json.dumps({'rows': rows}).encode()
            with self.assertRaises(urllib.error.HTTPError) as caught:
                opener.open(urllib.request.Request(route, data=body, headers={'Content-Type':'application/json'}))
            self.assertEqual(caught.exception.code, 403)
            with opener.open(urllib.request.Request(route, data=body, headers={'Content-Type':'application/json','X-Workbench-Key':self.store.key})) as response:
                self.assertEqual(json.load(response)['status'], 'done')
            doc = self.store.documents()[0]['statistics']
            with opener.open(base + '/api/document/' + doc['id']) as response:
                self.assertEqual(response.read(), Path(doc['path']).read_bytes())
            with opener.open(base + '/statistics.js') as response:
                self.assertEqual(response.status, 200)
        finally:
            http.shutdown(); http.server_close(); thread.join(timeout=3)


if __name__ == '__main__':
    unittest.main(verbosity=2)
