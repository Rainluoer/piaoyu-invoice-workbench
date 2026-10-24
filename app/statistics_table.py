"""从入库单的同一份商品数据生成统计表，不检索或推测安全等级及网址。"""
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import urlsplit
from invoice_company import company_needs_repair

CATEGORIES = ('试剂', '仪器', '耗材')
FIELDS = ('category', 'name', 'purpose', 'company', 'url', 'price', 'safety', 'notes')
HEADERS = ('类别', '名称', '用途', '公司', '网址', '价格', '安全等级', '备注')
TEMPLATE = Path(__file__).resolve().parent / 'assets' / '统计表格模板.docx'


def category_for(item):
    category = str(item.get('category') or '')
    if category in CATEGORIES:
        return category
    text = category + ' ' + str(item.get('product') or '')
    if any(word in text for word in ('试剂', '培养基', '缓冲液', '乙醇', '氯化钠', '检测试剂盒')):
        return '试剂'
    if any(word in text for word in ('仪器', '显微镜', '天平', '离心机', '光度计', '移液器', '示波器', '万用表')):
        return '仪器'
    return '耗材'


def money(value):
    if value is None or value == '':
        return ''
    number = Decimal(str(value))
    if not number.is_finite() or number < 0:
        raise ValueError('统计表价格无效')
    return format(number.quantize(Decimal('.01'), rounding=ROUND_HALF_UP), '.2f')


def rows_from_invoices(invoices, sellers=None):
    sellers = sellers or {}
    rows = []
    for index, inv in enumerate(invoices, 1):
        items = inv.get('items') or []
        for number, item in enumerate(items, 1):
            amount = inv.get('total') if len(items) == 1 else None
            if amount is None and item.get('amount') is not None:
                amount = Decimal(str(item['amount'])) + Decimal(str(item.get('tax') or 0))
            rows.append(dict(key=f'{index}:{number}', category=category_for(item),
                name=' '.join(str(item.get(field) or '').strip() for field in ('product', 'spec')).strip(),
                purpose='', company=str(sellers.get(inv.get('original_filename') or f'{index:03d}.pdf', '') or ''),
                url='', price=money(amount), safety='', notes=''))
    if not rows:
        raise ValueError('本次任务没有可用于统计的商品数据')
    return rows


def validate_rows(submitted, expected):
    if not isinstance(submitted, list) or len(submitted) != len(expected):
        raise ValueError('统计表商品数量与入库单不一致')
    by_key = {}
    for row in submitted:
        if not isinstance(row, dict) or not isinstance(row.get('key'), str) or row['key'] in by_key:
            raise ValueError('统计表商品清单无效')
        by_key[row['key']] = row
    if set(by_key) != {row['key'] for row in expected}:
        raise ValueError('统计表商品与入库单不一致')
    result = []
    for original in expected:
        row = by_key[original['key']]
        clean = {**original}
        for field, limit in (('category', 10), ('purpose', 300), ('company', 200), ('url', 2000), ('safety', 100), ('notes', 300)):
            value = row.get(field, '')
            if not isinstance(value, str) or len(value) > limit or any(ord(c) < 32 for c in value):
                raise ValueError('统计表字段格式无效或内容过长')
            clean[field] = value.strip()
        if clean['category'] not in CATEGORIES:
            raise ValueError('统计表类别只能选择试剂、仪器或耗材')
        if clean['url']:
            try:
                parsed = urlsplit(clean['url'])
                if parsed.scheme not in ('https', 'http') or not parsed.hostname or parsed.username or parsed.password or any(c.isspace() for c in clean['url']):
                    raise ValueError()
                parsed.port
            except ValueError:
                raise ValueError('网址需为完整的 http 或 https 地址') from None
        if clean['category'] != '试剂':
            clean['safety'] = ''
        if company_needs_repair(clean['company']):
            if not original.get('company') or company_needs_repair(original['company']):
                raise ValueError('公司名称混入票面标签或水印，请核对后填写')
            clean['company'] = original['company']
        if not clean['company']:
            clean['company'] = original.get('company', '')
        # 名称与含税价格来自同一次入库数据，不接受客户端改写造成两表不一致。
        result.append(clean)
    return result


def write_statistics(output, rows, template=TEMPLATE):
    # 基础台账仍可在缺少 Word 生成依赖时启动，由入库单可用性检查提示。
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    if not rows or any(row.get('category') not in CATEGORIES for row in rows):
        raise ValueError('统计表需包含有效商品和固定类别')
    document = Document(template)
    table = document.tables[0]
    if tuple(cell.text for cell in table.rows[0].cells) != HEADERS:
        raise ValueError('统计表模板列结构不匹配')
    sample = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for record in rows:
        table._tbl.append(deepcopy(sample))
        for field, cell in zip(FIELDS, table.rows[-1].cells):
            value = record.get(field, '')
            if field == 'safety' and record['category'] != '试剂':
                value = ''
            cell.text = value
    for section in document.sections:
        section.left_margin = section.right_margin = Cm(2)
    widths = (1.5, 3.6, 2.1, 2.8, 2.2, 1.7, 1.5, 1.6)
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        for height in list(properties.findall(qn('w:trHeight'))):
            properties.remove(height)
        if index == 0 and properties.find(qn('w:tblHeader')) is None:
            properties.append(OxmlElement('w:tblHeader'))
        if index > 0 and properties.find(qn('w:cantSplit')) is None:
            properties.append(OxmlElement('w:cantSplit'))
        for col, (cell, width) in enumerate(zip(row.cells, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or col in (0, 5, 6) else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = '宋体'
                    run.font.size = Pt(11)
                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    # 模板中表格后的空段不承载说明；不添加合计、来源、提示或安全性注释。
    after = False
    for element in list(document.element.body):
        if element is table._tbl:
            after = True
        elif after and element.tag != qn('w:sectPr'):
            document.element.body.remove(element)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.save(output)
