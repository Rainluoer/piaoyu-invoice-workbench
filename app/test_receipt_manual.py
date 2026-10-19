"""提取失败补填、草稿恢复、金额边界及真实本地模板续跑。"""
import copy
from contextlib import closing
import hashlib
import json
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import threading
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request

from docx import Document
from server import Store, LocalServer, Handler


class ReceiptManualTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        folder = self.root / '工具'
        folder.mkdir()
        self.original = folder / '扫描票.pdf'
        self.original.write_bytes(b'%PDF manual fixture')
        self.store = Store(self.root, self.root / 'app/data')
        self.store.scan()
        self.row = self.store.state()['invoices'][0]

    def tearDown(self):
        self.store.db.close()
        self.temp.cleanup()

    def failed_job(self, extra=False, command_failure=False):
        if extra:
            self.original.with_name('第二张.pdf').write_bytes(b'%PDF second fixture')
            self.store.scan()
        ids = [row['id'] for row in self.store.state()['invoices']]
        def extract(command, **kwargs):
            data = Path(command[command.index('--output') + 1])
            data.parent.mkdir(parents=True, exist_ok=True)
            data.write_text(json.dumps([{'original_filename':'001.pdf', 'invoice_no':'AUTO-001', 'total':12.34}]), encoding='utf-8')
            return SimpleNamespace(returncode=1 if command_failure else 0, stdout='', stderr='extract error' if command_failure else '')
        with patch('receipt_jobs.subprocess.run', side_effect=extract):
            job = self.store.receipts.create(ids, background=False)
        self.assertEqual(job['status'], 'failed')
        self.assertTrue(job['manual_available'])
        return job

    def completed_fields(self, job):
        invoices = self.store.receipts.manual_data(job['id'])['invoices']
        for index, inv in enumerate(invoices):
            inv.update(invoice_no=f'MANUAL-{index+1}', total='12.34', items=[dict(
                category='五金', product='连接件', spec='M5', unit='个', quantity='2',
                unit_price='', amount='10.00', tax='2.34')])
        return invoices

    def test_partial_results_are_matched_by_filename_and_draft_survives_restart(self):
        job = self.failed_job(extra=True)
        rows = self.store.receipts.manual_data(job['id'])['invoices']
        self.assertEqual(rows[0]['invoice_no'], 'AUTO-001')
        self.assertEqual(rows[0]['total'], '12.34')
        self.assertEqual(rows[1]['invoice_no'], '')
        self.assertTrue(all(not row['extracted_complete'] for row in rows))
        rows[1]['items'][0]['product'] = '尚未填完的商品'
        self.store.receipts.save_manual(job['id'], rows)
        self.store.db.close()
        self.store = Store(self.root, self.root / 'app/data')
        restored = self.store.receipts.manual_data(job['id'])['invoices']
        self.assertEqual(restored[1]['items'][0]['product'], '尚未填完的商品')
        self.assertTrue(self.store.receipts.get(job['id'])['manual_available'])
        with closing(sqlite3.connect(self.store.backup())) as db:
            payload = json.loads(db.execute('SELECT payload FROM receipt_jobs WHERE id=?', (job['id'],)).fetchone()[0])
        self.assertEqual(payload['manual_draft'], restored)

    def test_real_rename_and_fill_resume_without_reextracting_or_changing_source(self):
        job = self.failed_job(command_failure=True)
        invoices = self.completed_fields(job)
        invoices[0]['items'].append(dict(category='五金', product='螺母', spec='', unit='个', quantity='1', unit_price='', amount='3.00', tax=''))
        invoices[0]['total'] = '15.34'
        original_hash = hashlib.sha256(self.original.read_bytes()).hexdigest()
        extracted_before = Path(job['data_path']).read_bytes()
        import subprocess
        with patch('receipt_jobs.subprocess.run', wraps=subprocess.run) as runner:
            done = self.store.receipts.save_manual(job['id'], invoices, resume=True, background=False)
        self.assertEqual(done['status'], 'done', done['error'])
        self.assertEqual([Path(call.args[0][1]).name for call in runner.call_args_list], ['rename_files.py','fill_receipt.py'])
        self.assertFalse(done['manual_available'])
        self.assertEqual(done['total_cents'], 1534)
        self.assertEqual(done['item_count'], 2)
        data = Path(done['data_path'])
        self.assertEqual(data.with_name('invoice_data_extracted.json').read_bytes(), extracted_before)
        normalized = json.loads(data.read_text(encoding='utf-8'))
        self.assertEqual(normalized[0]['items'][0]['unit_price'], 5)
        document = Document(done['output'])
        content = '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn('连接件', content)
        self.assertIn('螺母', content)
        self.assertIn('15.34', content)
        self.assertEqual(hashlib.sha256(self.original.read_bytes()).hexdigest(), original_hash)
        self.assertEqual(self.store.state()['invoices'][0]['path'], str(self.original))
        self.assertEqual(len(self.store.documents()), 1)

    def test_invalid_entered_numbers_and_path_characters_cannot_resume(self):
        job = self.failed_job()
        valid = self.completed_fields(job)
        cases = [('total','NaN'), ('total','Infinity'), ('total','-1'), ('total','12.345')]
        for field, value in cases:
            submitted = copy.deepcopy(valid)
            submitted[0][field] = value
            with self.subTest(field=field,value=value), self.assertRaises(ValueError):
                self.store.receipts.save_manual(job['id'], submitted, resume=True, background=False)
            self.assertFalse(self.store.receipts.active.locked())
        for field, value in [('quantity','0'),('quantity','NaN'),('product','../越界'),('tax','-1')]:
            submitted = copy.deepcopy(valid)
            submitted[0]['items'][0][field] = value
            with self.subTest(field=field,value=value), self.assertRaises(ValueError):
                self.store.receipts.save_manual(job['id'], submitted, resume=True, background=False)
        self.assertFalse(Path(job['output']).is_file())
        self.assertEqual(self.store.receipts.get(job['id'])['status'], 'failed')

    def test_high_precision_unit_prices_are_preserved_and_real_template_generates(self):
        job=self.failed_job()
        invoices=self.completed_fields(job)
        invoices[0]['items'][0]['unit_price']='35.2920353982'
        # 单商品采用票面合计，不能用单价舍入或识别明细反向否决票面合计。
        invoices[0]['total']='12.32'
        for price in ['35.2920353982','15.6902654867','155.3398058252427']:
            invoices[0]['items'][0]['unit_price']=price
            draft,data=self.store.receipts._manual_invoices(job,invoices,True)
            self.assertEqual(draft[0]['items'][0]['unit_price'],price)
            self.assertEqual(data[0]['items'][0]['unit_price'],float(price))
        done=self.store.receipts.save_manual(job['id'],invoices,resume=True,background=False)
        self.assertEqual(done['status'],'done',done['error'])
        data=json.loads(Path(done['data_path']).read_text(encoding='utf-8'))
        self.assertEqual(data[0]['items'][0]['unit_price'],155.3398058252427)
        row=Document(done['output']).tables[1].rows[1]
        self.assertEqual(row.cells[6].text,'155.34')
        self.assertEqual(row.cells[7].text,'12.32')

    def test_failed_manual_audit_rolls_back_status_and_draft_then_can_retry(self):
        job = self.failed_job()
        invoices = self.completed_fields(job)
        before = self.store.db.execute('SELECT payload FROM receipt_jobs WHERE id=?', (job['id'],)).fetchone()[0]
        self.store.db.execute("CREATE TRIGGER qa_reject_manual BEFORE INSERT ON events WHEN NEW.action='补填入库单' BEGIN SELECT RAISE(ABORT,'audit rejected'); END")
        self.store.db.commit()
        with self.assertRaises(sqlite3.IntegrityError):
            self.store.receipts.save_manual(job['id'], invoices, resume=True)
        self.assertFalse(self.store.receipts.active.locked())
        self.assertFalse(self.store.db.in_transaction)
        self.assertEqual(self.store.db.execute('SELECT payload FROM receipt_jobs WHERE id=?', (job['id'],)).fetchone()[0], before)
        self.assertTrue(self.store.receipts.get(job['id'])['manual_available'])
        self.store.db.execute('DROP TRIGGER qa_reject_manual')
        self.store.db.commit()
        done = self.store.receipts.save_manual(job['id'], invoices, resume=True, background=False)
        self.assertEqual(done['status'], 'done', done['error'])

    def test_failed_manual_worker_launch_preserves_draft_and_can_retry(self):
        job = self.failed_job()
        invoices = self.completed_fields(job)
        with patch('receipt_jobs.threading.Thread.start', side_effect=RuntimeError('worker unavailable')):
            failed = self.store.receipts.save_manual(job['id'], invoices, resume=True)
        self.assertEqual(failed['status'], 'failed')
        self.assertIn('无法启动生成任务', failed['error'])
        self.assertTrue(failed['manual_available'])
        self.assertFalse(self.store.receipts.active.locked())
        self.assertEqual(self.store.receipts.manual_data(job['id'])['invoices'][0]['items'][0]['spec'], 'M5')
        self.store.receipts.save_manual(job['id'], invoices, resume=True)
        self.store.receipts.thread.join(timeout=5)
        self.assertFalse(self.store.receipts.thread.is_alive())
        done = self.store.receipts.get(job['id'])
        self.assertEqual(done['status'], 'done', done['error'])
        self.assertFalse(self.store.receipts.active.locked())

    def test_failed_new_worker_launch_does_not_leave_a_running_job(self):
        with patch('receipt_jobs.threading.Thread.start', side_effect=RuntimeError('worker unavailable')):
            failed = self.store.receipts.create([self.row['id']])
        self.assertEqual(failed['status'], 'failed')
        self.assertIn('无法启动生成任务', failed['error'])
        self.assertFalse(self.store.receipts.active.locked())
        self.assertFalse(Path(failed['output']).exists())
        retried = self.failed_job()
        self.assertNotEqual(retried['id'], failed['id'])

    def test_missing_tax_does_not_block_but_complete_multi_item_mismatch_is_detected(self):
        job=self.failed_job()
        invoices=self.completed_fields(job)
        invoices[0]['items'].append({'product':'第二项','amount':'10.00','tax':'0'})
        with self.assertRaisesRegex(ValueError,'不等于价税合计'):
            self.store.receipts._manual_invoices(job,invoices,True)
        invoices[0]['items'][1]['tax']=''
        self.store.receipts._manual_invoices(job,invoices,True)
        for invalid in ['NaN','Infinity','-1','￥35.29']:
            invoices[0]['items'][0]['unit_price']=invalid
            with self.subTest(invalid=invalid),self.assertRaises(ValueError):
                self.store.receipts._manual_invoices(job,invoices,True)

    def test_all_fields_blank_can_generate_with_real_blank_cells_and_totals(self):
        job = self.failed_job(extra=True)
        invoices = self.store.receipts.manual_data(job['id'])['invoices']
        for inv in invoices:
            inv.update(invoice_no='',total='',items=[{}])
        done = self.store.receipts.save_manual(job['id'],invoices,resume=True,background=False)
        self.assertEqual(done['status'],'done',done['error'])
        self.assertIsNone(done['total_cents'])
        self.assertEqual(done['chinese_total'],'')
        data=json.loads(Path(done['data_path']).read_text(encoding='utf-8'))
        self.assertTrue(all(inv['total'] is None and inv['invoice_no']=='' for inv in data))
        document=Document(done['output'])
        table=document.tables[1]
        for row in table.rows[1:3]:
            self.assertTrue(all(cell.text=='' for cell in row.cells[2:8]))
        self.assertEqual(table.rows[8].cells[7].text,'')
        self.assertEqual(table.rows[8].cells[0].text,'金额合计（大写）：')
        self.assertEqual(len(document.tables[0].rows),3)
        self.assertEqual(len(self.store.documents()),1)

    def test_only_product_and_known_total_are_enough_and_missing_fields_stay_blank(self):
        job=self.failed_job()
        invoices=self.store.receipts.manual_data(job['id'])['invoices']
        invoices[0].update(invoice_no='',items=[{'product':'手动补填品名'}])
        done=self.store.receipts.save_manual(job['id'],invoices,resume=True,background=False)
        self.assertEqual(done['status'],'done',done['error'])
        self.assertEqual(done['total_cents'],1234)
        row=Document(done['output']).tables[1].rows[1]
        self.assertEqual('手动补填品名',row.cells[2].text)
        self.assertTrue(all(cell.text=='' for cell in row.cells[4:7]))
        self.assertEqual(row.cells[7].text,'12.34')

    def test_multiple_missing_invoice_numbers_do_not_deduplicate_totals_and_zero_is_known(self):
        job=self.failed_job(extra=True)
        invoices=self.store.receipts.manual_data(job['id'])['invoices']
        for inv,total in zip(invoices,['0','10.00']):
            inv.update(invoice_no='',total=total,items=[{}])
        done=self.store.receipts.save_manual(job['id'],invoices,resume=True,background=False)
        self.assertEqual(done['status'],'done',done['error'])
        table=Document(done['output']).tables[1]
        self.assertEqual(table.rows[1].cells[7].text,'0.00')
        self.assertEqual(table.rows[2].cells[7].text,'10.00')
        self.assertEqual(table.rows[8].cells[7].text,'10.00')
        self.assertEqual(done['total_cents'],1000)

    def test_wrong_sources_duplicate_invoice_numbers_and_changed_copy_are_rejected(self):
        job = self.failed_job(extra=True)
        valid = self.completed_fields(job)
        wrong = copy.deepcopy(valid)
        wrong[1]['source_id'] = wrong[0]['source_id']
        with self.assertRaises(ValueError):
            self.store.receipts.save_manual(job['id'], wrong, resume=True, background=False)
        wrong = copy.deepcopy(valid)
        wrong[1]['invoice_no'] = wrong[0]['invoice_no']
        with self.assertRaises(ValueError):
            self.store.receipts.save_manual(job['id'], wrong, resume=True, background=False)
        (Path(job['output']).parent / '001.pdf').write_bytes(b'changed copy')
        with self.assertRaisesRegex(ValueError, '副本已变化'):
            self.store.receipts.save_manual(job['id'], valid, resume=True, background=False)
        self.assertFalse(self.store.receipts.active.locked())

    def test_busy_or_later_step_failure_cannot_be_manually_overwritten(self):
        job = self.failed_job()
        fields = self.completed_fields(job)
        self.store.receipts.active.acquire()
        try:
            with self.assertRaisesRegex(ValueError, '正在生成'):
                self.store.receipts.save_manual(job['id'], fields)
        finally:
            self.store.receipts.active.release()
        job['stage'] = 2
        self.store.receipts._save(job)
        with self.assertRaisesRegex(ValueError, '不能补填'):
            self.store.receipts.manual_data(job['id'])

    def test_http_manual_form_data_and_authenticated_draft_save(self):
        job = self.failed_job()
        server = LocalServer(('127.0.0.1',0), Handler)
        server.store = self.store
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))
        url = f'http://127.0.0.1:{server.server_port}/api/receipt-jobs/{job["id"]}/manual'
        try:
            with opener.open(url) as response:
                payload = json.load(response)
            self.assertEqual(payload['job_id'], job['id'])
            body = json.dumps({'invoices':payload['invoices']}).encode()
            with self.assertRaises(urllib.error.HTTPError) as error:
                opener.open(urllib.request.Request(url,data=body,headers={'Content-Type':'application/json'}))
            self.assertEqual(error.exception.code, 403)
            with opener.open(urllib.request.Request(url,data=body,headers={'Content-Type':'application/json','X-Workbench-Key':self.store.key})) as response:
                saved = json.load(response)
            self.assertTrue(saved['manual_saved_at'])
            self.assertTrue(saved['manual_available'])
        finally:
            server.shutdown(); server.server_close(); thread.join(timeout=3)


if __name__ == '__main__':
    unittest.main(verbosity=2)
